"""Export one content source to XeLaTeX and DOCX, with claim/source hashes."""
from pathlib import Path
import csv, hashlib, json, re, sys, subprocess, shutil
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from manuscript import TITLE, SUBTITLE, BLOCKS, REFERENCES, SOURCES

HERE=Path(__file__).resolve().parent
ROOT=HERE.parent
REPOSITORY_URL='https://github.com/lliooly/laya-imdb'
MODEL_HUB_URL='https://huggingface.co/shishishi3/laya-deberta-imdb-ensemble'
metadata_path=HERE/'submission.json'
metadata=json.loads(metadata_path.read_text())
required_identity=('name','student_id','class_name')
if any(not isinstance(metadata.get(k),str) or not metadata[k].strip()
       or re.search(r'待填写|TODO|TBD',metadata[k],re.I) for k in required_identity):
    raise ValueError('Complete name, student_id and class_name in submission.json before exporting.')
identity={k:metadata[k].strip() for k in required_identity}
identity_lines=[f'姓名：{identity["name"]}    学号：{identity["student_id"]}',
                f'班级：{identity["class_name"]}']
build_time=datetime.now(ZoneInfo('Asia/Shanghai'))
author=identity['name']
def esc(s):
    value=''.join({'\\':r'\textbackslash{}','&':r'\&','%':r'\%','$':r'\$','#':r'\#','_':r'\_','{':r'\{','}':r'\}'}.get(c,c) for c in str(s))
    for a,b in [('×10⁻²⁰',r'\ensuremath{\times10^{-20}}'),('×10⁻⁵',r'\ensuremath{\times10^{-5}}'),('λ',r'\ensuremath{\lambda}'),('ℓ',r'\ensuremath{\ell}'),('≈',r'\ensuremath{\approx}'),('≤',r'\ensuremath{\le}')]:value=value.replace(a,b)
    return value

def tex_prose(text):
    # Substitute only explicit numbered citations and prose references; equations remain verbatim.
    pattern=r' \[([1-9]\d*)\]|图 (\d+)|表 (\d+)|式（(\d+)）'
    result=[];last=0
    for m in re.finditer(pattern,text):
        result.append(esc(text[last:m.start()]))
        cite,fig,tab,eq=m.groups()
        if cite:result.append(r'\cite{'+REFERENCES[int(cite)-1][0]+'}')
        elif fig:result.append(r'图~\ref{fig:f'+fig+'}')
        elif tab:result.append(r'表~\ref{tab:t'+tab+'}')
        else:result.append(r'式~\eqref{eq:e'+eq+'}')
        last=m.end()
    result.append(esc(text[last:]));return ''.join(result)

tex=[r'''% Generated from manuscript.py by build_report.py. UTF-8; compile with XeLaTeX.
\documentclass[UTF8,fontset=fandol,a4paper,11pt]{ctexart}
\usepackage[margin=23mm,top=22mm,bottom=22mm]{geometry}
\usepackage{amsmath,amssymb,graphicx,booktabs,tabularx,array,float,adjustbox}
\usepackage[font=small,labelfont=bf]{caption}
\usepackage{xcolor}
\usepackage[hidelinks,unicode]{hyperref}
\usepackage{microtype}
\usepackage{setspace}
\usepackage[section]{placeins}
\usepackage{flafter}
\captionsetup{labelsep=space}
\setstretch{1.18}
\setlength{\parskip}{3pt}
\setlength{\emergencystretch}{3em}
\setlength{\textfloatsep}{12pt}
\raggedbottom
\pagestyle{plain}
\ctexset{section={format=\Large\bfseries},subsection={format=\large\bfseries}}
\hypersetup{pdftitle={IMDB 情感分类项目报告},pdfauthor={'''+esc(author)+r'''},pdfsubject={'''+esc(SUBTITLE)+r'''},pdfkeywords={IMDB, Laya, DeBERTa-v3, 情感分类, 概率融合}}
\newcolumntype{Y}{>{\raggedright\arraybackslash}X}
\begin{document}
\begin{center}
{\LARGE\bfseries '''+esc(TITLE)+r'''}\par
\vspace{7pt}
{\normalsize '''+esc(SUBTITLE)+r'''}\par
\vspace{5pt}
{\small 2026 年 10 月}\par
\vspace{5pt}
{\small '''+esc(identity_lines[0]).replace('    ',r'\qquad ')+r'''}\par
{\small '''+esc(identity_lines[1]).replace('    ',r'\qquad ')+r'''}\par
{\small 项目源码：\url{'''+REPOSITORY_URL+r'''}}\par
{\small 模型权重：\url{'''+MODEL_HUB_URL+r'''}}
\end{center}
\vspace{4pt}
''']
tex_fig=tex_tab=tex_eq=0
for i,b in enumerate(BLOCKS):
    if b.get('source'):tex.append('% Evidence: '+b['source'])
    kind=b['kind']
    if kind=='h':
        cmd={0:'section*',1:'section',2:'subsection'}[b['level']]
        tex.append('\\'+cmd+'{'+esc(b['text'])+'}')
    elif kind=='p':tex.append(tex_prose(b['text'])+'\n')
    elif kind=='equation':
        tex_eq+=1
        tex.append(r'\begin{equation}\label{eq:e'+str(tex_eq)+r'}\begin{adjustbox}{max width=.94\linewidth}$\displaystyle '+b['tex']+r'$\end{adjustbox}\end{equation}')
    elif kind=='figure':
        tex_fig+=1
        tex.append(r'\begin{figure}[htbp]\centering\includegraphics[width=\linewidth]{figures/'+b['name']+r'.pdf}\caption{'+esc(b['caption'])+r'}\label{fig:f'+str(tex_fig)+r'}\end{figure}')
    elif kind=='table':
        cols=len(b['headers'])
        tex_tab+=1
        placement='H' if tex_tab==3 else 'htbp'
        tex.append(r'\begin{table}['+placement+r']\centering\caption{'+esc(b['caption'])+r'}\label{tab:t'+str(tex_tab)+r'}\small\renewcommand{\arraystretch}{1.25}')
        if tex_tab==3:
            tex.append(r'\footnotesize\setlength{\tabcolsep}{3pt}')
            layout=r'@{\extracolsep{\fill}}l l r r r r r@{}'
            table_env='tabular*'
        else:layout='Y'*cols;table_env='tabularx'
        tex.append(r'\begin{'+table_env+r'}{\linewidth}{'+layout+r'}\toprule')
        tex.append(' & '.join(r'\textbf{'+esc(s)+'}' for s in b['headers'])+r'\\\midrule')
        tex += [' & '.join(esc(s) for s in row)+r'\\' for row in b['rows']]
        tex.append(r'\bottomrule\end{'+table_env+r'}\end{table}')
tex.append(r'\FloatBarrier\begin{thebibliography}{99}')
for key,ref,url in REFERENCES:tex.append(r'\bibitem{'+key+'} '+esc(ref)+r' \url{'+url+'}')
tex.append(r'\end{thebibliography}\end{document}')
(HERE/'main.tex').write_text('\n'.join(tex)+'\n')

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
doc=Document();sec=doc.sections[0]
for el in doc.styles.element.xpath('.//w:pBdr'):
    el.getparent().remove(el)
for el in doc.element.xpath('.//w:pBdr | .//w:docGrid'):
    el.getparent().remove(el)
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=sec.bottom_margin=Cm(2.1);sec.left_margin=sec.right_margin=Cm(2.3)
def fonts(style,east='SimSun',western='Times New Roman',size=11):
    style.font.name=western;style.font.size=Pt(size)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),east)
fonts(doc.styles['Normal']);normal=doc.styles['Normal'].paragraph_format
normal.line_spacing=1.15;normal.space_after=Pt(4);normal.first_line_indent=Cm(.74)
normal.widow_control=True
snap=OxmlElement('w:snapToGrid');snap.set(qn('w:val'),'0');doc.styles['Normal'].element.get_or_add_pPr().append(snap)
for name,size in [('Title',19),('Heading 1',15),('Heading 2',12),('Caption',9)]:
    fonts(doc.styles[name],east='SimHei' if name!='Caption' else 'SimSun',size=size)
    doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    if name=='Caption':doc.styles[name].font.bold=False
    pf=doc.styles[name].paragraph_format;pf.first_line_indent=Pt(0);pf.space_before=Pt(10);pf.space_after=Pt(6)
    if name!='Caption':pf.keep_with_next=True
p=doc.add_paragraph(style='Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run(TITLE)
p=doc.add_paragraph(SUBTITLE);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
p=doc.add_paragraph('2026 年 10 月');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
for run in p.runs:run.font.size=Pt(10)
for text in identity_lines:
    p=doc.add_paragraph(text);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(2)
    for run in p.runs:run.font.size=Pt(9)
for label,url in [('项目源码：',REPOSITORY_URL),('模型权重：',MODEL_HUB_URL)]:
    p=doc.add_paragraph(label);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(2)
    for run in p.runs:run.font.size=Pt(9)
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'18');pr.append(size);r.append(pr)
    t=OxmlElement('w:t');t.text=url;r.append(t);link.append(r);p._p.append(link)
fig_n=tab_n=eq_n=section_n=sub_n=0
for b in BLOCKS:
    kind=b['kind']
    if kind=='h':
        lv=b['level'];text=b['text']
        if lv==1:section_n+=1;sub_n=0;text=f'{section_n}  {text}'
        elif lv==2:sub_n+=1;text=f'{section_n}.{sub_n}  {text}'
        doc.add_heading(text,level=1 if lv in [0,1] else 2)
    elif kind=='p':doc.add_paragraph(b['text'])
    elif kind=='equation':
        eq_n+=1;p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(8)
        # Render the exact TeX equation for portable Word typography; editable source remains in main.tex.
        folder=HERE/'qa/equations';folder.mkdir(parents=True,exist_ok=True)
        stem=folder/f'equation-{eq_n:02d}'
        equation_source=(r'\documentclass[11pt,border=3pt]{standalone}'+'\n'+
            r'\usepackage{amsmath,amssymb}\begin{document}$\displaystyle '+b['tex']+
            r'\qquad\text{('+str(eq_n)+r')}$\end{document}')
        stem.with_suffix('.tex').write_text(equation_source)
        result=subprocess.run([shutil.which('xelatex') or '/Library/TeX/texbin/xelatex','-interaction=nonstopmode','-halt-on-error',
            '-output-directory='+str(folder),str(stem.with_suffix('.tex'))],cwd=ROOT,capture_output=True,text=True)
        if result.returncode:raise RuntimeError(result.stdout[-2500:])
        subprocess.run([shutil.which('pdftoppm') or '/opt/homebrew/bin/pdftoppm','-r','300','-singlefile','-png',str(stem.with_suffix('.pdf')),str(stem)],check=True,capture_output=True)
        from PIL import Image
        im=Image.open(stem.with_suffix('.png'));width=min(16.2,im.width/300*2.54)
        p.add_run().add_picture(str(stem.with_suffix('.png')),width=Cm(width))
        p._p.xpath('.//wp:docPr')[0].set('descr',b['word'])
    elif kind=='figure':
        fig_n+=1;p=doc.add_paragraph();p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
        p.paragraph_format.line_spacing=1;p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(0)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(HERE/'figures'/(b['name']+'.png')),width=Cm(15.5))
        inline=p._p.xpath('.//wp:docPr')[0];inline.set('descr',b['caption'])
        c=doc.add_paragraph(f'图 {fig_n}  '+b['caption'],style='Caption');c.paragraph_format.keep_with_next=False
    elif kind=='table':
        tab_n+=1;c=doc.add_paragraph(f'表 {tab_n}  '+b['caption'],style='Caption');c.paragraph_format.keep_with_next=True
        table=doc.add_table(rows=1,cols=len(b['headers']));table.alignment=WD_TABLE_ALIGNMENT.CENTER
        table.autofit=False
        width=16.4/len(b['headers'])
        widths=[.8,3.0,2.6,2.0,2.4,2.4,2.1] if tab_n==3 else [width]*len(b['headers'])
        widths=[v*16.4/sum(widths) for v in widths]
        for col,w in zip(table.columns,widths):col.width=Cm(w)
        for cell,text in zip(table.rows[0].cells,b['headers']):cell.text=text
        pr=table.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');pr.append(rep)
        for row in b['rows']:
            for cell,text in zip(table.add_row().cells,row):cell.text=str(text)
        for j,row in enumerate(table.rows):
            trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
            for cell,w in zip(row.cells,widths):
                cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr=cell._tc.get_or_add_tcPr();marg=OxmlElement('w:tcMar')
                for side in ['top','bottom','left','right']:
                    padding='30' if tab_n in [2,3] and side in ['top','bottom'] else '65'
                    e=OxmlElement('w:'+side);e.set(qn('w:w'),padding);e.set(qn('w:type'),'dxa');marg.append(e)
                tcpr.append(marg)
                borders=OxmlElement('w:tcBorders')
                if j==0 or j==len(table.rows)-1:
                    e=OxmlElement('w:bottom');e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'808080');borders.append(e)
                tcpr.append(borders)
                if j==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F0F3F5');tcpr.append(shade)
                for par in cell.paragraphs:
                    par.paragraph_format.first_line_indent=Pt(0);par.paragraph_format.line_spacing=1.1;par.paragraph_format.space_after=Pt(2)
                    if tab_n==3:par.paragraph_format.line_spacing=1;par.paragraph_format.space_after=Pt(0)
                    par.paragraph_format.keep_with_next=(j<len(table.rows)-1 if tab_n in [2,3] else j==0)
                    for r in par.runs:r.font.size=Pt(8.5);r.bold=j==0
        doc.add_paragraph().paragraph_format.space_after=Pt(0)
doc.add_heading('参考文献',level=1)
for i,(_,ref,url) in enumerate(REFERENCES,1):
    p=doc.add_paragraph(f'[{i}] {ref} ')
    p.paragraph_format.left_indent=Cm(.74);p.paragraph_format.first_line_indent=Cm(-.74)
    p.paragraph_format.line_spacing=1.1
    p.paragraph_format.keep_with_next=i<len(REFERENCES)
    for r in p.runs:r.font.size=Pt(10)
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    link=OxmlElement('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');size=OxmlElement('w:sz');size.set(qn('w:val'),'20');pr.append(size);r.append(pr)
    t=OxmlElement('w:t');t.text=url;r.append(t);link.append(r);p._p.append(link)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title=TITLE;doc.core_properties.subject=SUBTITLE;doc.core_properties.author=author;doc.core_properties.last_modified_by=author
doc.core_properties.created=doc.core_properties.modified=build_time.astimezone(timezone.utc)
doc.core_properties.keywords='IMDB; Laya; DeBERTa-v3; 情感分类; 概率融合'
doc.core_properties.comments=''
from word_typography import normalize_word_fonts
normalize_word_fonts(doc)
doc.save(HERE/'IMDB项目报告.docx')

source_manifest={}
for sid,rel in SOURCES.items():
    p=ROOT/rel
    if not p.exists():raise FileNotFoundError(p)
    source_manifest[sid]={'path':rel,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),
       'verification':'opened and checked by assistant; human review not recorded'}
for key,ref,url in REFERENCES:source_manifest['ref_'+key]={'reference':ref,'url':url,'access_date':'2026-10-19','verification':'primary page opened by assistant; human review not recorded'}
(HERE/'analysis/source_manifest.json').write_text(json.dumps(source_manifest,indent=2,ensure_ascii=False)+'\n')
with (HERE/'analysis/claims.csv').open('w') as f:
    w=csv.writer(f);w.writerow(['block_id','kind','content_sha256','evidence_ids'])
    for i,b in enumerate(BLOCKS):
        content=b.get('text',b.get('caption',b.get('tex','')))
        w.writerow([f'C{i+1:03}',b['kind'],hashlib.sha256(content.encode()).hexdigest(),b.get('source','')])
(HERE/'analysis/report_content.json').write_text(json.dumps({'title':TITLE,'blocks':BLOCKS,'references':REFERENCES},ensure_ascii=False,indent=2)+'\n')
print(json.dumps({'blocks':len(BLOCKS),'figures':fig_n,'tables':tab_n,'equations':eq_n,'sources':len(source_manifest)},ensure_ascii=False))
